import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Set page margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    # Add page border (basic single line border around the whole page)
    sectPr = section._sectPr
    pg_border_xml = f'''
    <w:pgBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="8" w:space="20" w:color="555555"/>
        <w:left w:val="single" w:sz="8" w:space="20" w:color="555555"/>
        <w:bottom w:val="single" w:sz="8" w:space="20" w:color="555555"/>
        <w:right w:val="single" w:sz="8" w:space="20" w:color="555555"/>
    </w:pgBorders>
    '''
    sectPr.append(parse_xml(pg_border_xml))

# Heading: "hii soham" in Bold and Font Size 25
title_p = doc.add_paragraph()
title_run = title_p.add_run("hii soham")
title_run.bold = True
title_run.font.size = Pt(25)
title_run.font.name = "Arial"
title_run.font.color.rgb = RGBColor(15, 23, 42)
title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
title_p.paragraph_format.space_after = Pt(12)

# Subtitle
sub_p = doc.add_paragraph()
sub_run = sub_p.add_run("Source Code (HTML & CSS):")
sub_run.bold = True
sub_run.font.size = Pt(11)
sub_run.font.name = "Arial"
sub_run.font.color.rgb = RGBColor(71, 85, 105)
sub_p.paragraph_format.space_after = Pt(8)

# Read all code from index.html
with open(r"d:\cloud pr\index.html", "r", encoding="utf-8") as f:
    code_content = f.read()

# Create a bordered table box for the code
table = doc.add_table(rows=1, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

cell = table.cell(0, 0)
cell.width = Inches(6.8)

# Single line border around the code box
tcPr = cell._tc.get_or_add_tcPr()
border_xml = f'''
<w:tcBorders {nsdecls("w")}>
    <w:top w:val="single" w:sz="6" w:color="94A3B8"/>
    <w:left w:val="single" w:sz="6" w:color="94A3B8"/>
    <w:bottom w:val="single" w:sz="6" w:color="94A3B8"/>
    <w:right w:val="single" w:sz="6" w:color="94A3B8"/>
</w:tcBorders>
'''
tcPr.append(parse_xml(border_xml))

# Shading for code box
shade_xml = f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'
tcPr.append(parse_xml(shade_xml))

# Add code into cell
p = cell.paragraphs[0]
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.line_spacing = 1.15

code_run = p.add_run(code_content)
code_run.font.name = "Consolas"
code_run.font.size = Pt(8.5)
code_run.font.color.rgb = RGBColor(30, 41, 59)

# Save document
output_path = r"d:\cloud pr\hii_soham.docx"
doc.save(output_path)
print(f"Successfully generated: {output_path}")
